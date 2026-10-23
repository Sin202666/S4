# -*- coding: utf-8 -*-
"""Generate ArmWrestle Excel config sheets from current design numbers."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

OUT = Path(__file__).resolve().parents[1] / "config"


def _sheet(wb: Workbook, headers_zh: list[str], names: list[str], types: list[str], rows: list[list]):
	ws = wb.active
	ws.append(headers_zh)
	ws.append(names)
	ws.append(types)
	for row in rows:
		ws.append(row)


def _save(wb: Workbook, name: str):
	path = OUT / name
	wb.save(path)
	print(f"wrote {path}")


def gen_const():
	wb = Workbook()
	# 备注列仅给人看；运行时 ArmWrestleConfig 只读 Value
	rows = [
		["EXIT_PUSH_STUDS", 3, "出场时沿 EXIT 朝向外推的距离（studs）"],
		["WIN_CINEMATIC_SEC", 4.6, "获胜演出总时长（秒）；服务端交还控制须大于它"],
		["WIN_RELEASE_DELAY_SEC", 5.2, "获胜后服务端交还控制/出场延迟（秒）"],
		["LOSE_RELEASE_DELAY_SEC", 2.6, "失败后服务端交还控制/出场延迟（秒）"],
		["LIMITED_ARM_ID", 12, "限量手臂 Id（审判之臂）"],
		["LIMITED_SOLD_DEFAULT", 96, "限量臂已售默认值（全服库存展示兜底）"],
		["NEAR_MACHINE_STUDS", 25, "靠近材质箱的判定距离（studs）"],
		["TABLE_PROMPT_DISTANCE", 18, "比赛桌扳手腕 E 交互距离（studs）"],
		["ROLL_ANIM_SEC", 3.2, "材质箱抽取 3D 表演总秒数；冷却与之对齐"],
		["FRIEND_BOOST_PER", 0.1, "同服每位好友现金加成（+10%）"],
		["FRIEND_BOOST_MAX", 0.5, "好友现金加成封顶（+50%）"],
		["PREMIUM_BOOST", 0.25, "Roblox 会员现金加成（+25%）"],
		["OFFLINE_RATE", 0.1, "离线力量折算：在线挂机效率的 10%"],
		["OFFLINE_MIN_SEC", 300, "离线最短门槛（秒）；线上 5 分钟，Studio 代码内改为 10 秒"],
		["OFFLINE_MAX_SEC", 28800, "离线累计封顶（秒，8 小时）"],
		["MENU_GIFT_CASH_MULT", 10, "菜单挽留礼物：已解锁关最强对手奖励 × 本倍数"],
		["ZoneLength", 170, "单关世界区 X 向长度（studs）"],
		["ZoneWidth", 130, "单关世界区 Z 向宽度（studs）"],
		["BattleTime", 15, "扳手腕对局时长（秒）"],
		["TapPower", 0.035, "单次点击对进度条贡献"],
		["OpponentDrift", 0.13, "对手持续回推强度；需持续点击才能压住"],
		["MaxTapsPerSecond", 14, "玩家点击速率上限（防脚本）"],
		["AutoTapsPerSecond", 8, "自动点击通行证每秒点击次数"],
	]
	_sheet(
		wb,
		["键", "数值", "备注"],
		["Key", "Value", "Remark"],
		["string", "number", "string"],
		rows,
	)
	_save(wb, "AWConst.xlsx")


def gen_material():
	wb = Workbook()
	# Id key; Color as R|G|B number[]; empty Price/ProductKey -> 0 / ""
	# 抽取权重在 AWLevel.PoolWeights（按箱），不挂材质行
	# 购买价挂材质行；抽取免费；奥术只走 Robux（Price=0）
	rows = [
		["Sand", "沙子", 1.2, 1, "Sand", 226, 205, 145, 40, "", "rbxassetid://122192010470765"],
		["Wood", "木头", 1.5, 1, "Wood", 160, 110, 60, 80, "", "rbxassetid://128126726660785"],
		["Grass", "草地", 2, 2, "Grass", 90, 170, 60, 150, "", "rbxassetid://126845964848626"],
		["Snow", "冰雪", 2.5, 2, "Snow", 235, 245, 255, 500, "", "rbxassetid://122005236129409"],
		["Rock", "岩石", 3.5, 3, "Slate", 110, 110, 120, 1500, "", "rbxassetid://130795413314674"],
		["Metal", "金属", 5, 3, "DiamondPlate", 170, 175, 185, 5000, "", "rbxassetid://126111622075252"],
		["Water", "流水", 7, 4, "Glass", 60, 150, 255, 15000, "", "rbxassetid://116870538790500"],
		["Brick", "砖块", 10, 4, "Brick", 170, 70, 50, 50000, "", "rbxassetid://124670707684711"],
		["Magma", "岩浆", 15, 4, "CrackedLava", 255, 100, 30, 200000, "", "rbxassetid://104796958921544"],
		["Neon", "霓虹", 25, 4, "Neon", 255, 60, 200, 800000, "", "rbxassetid://81720242495997"],
		["Diamond", "钻石", 40, 5, "Glass", 120, 240, 255, 2000000, "", "rbxassetid://100212317246111"],
		["Arcane", "奥术", 100, 5, "ForceField", 170, 80, 255, 0, "ArcaneMaterial", "rbxassetid://110621231829163"],
	]
	_sheet(
		wb,
		["材质Id", "中文名", "力量倍率", "稀有度档", "材质枚举", "色R", "色G", "色B", "购买价", "商品键", "图标"],
		["Id", "ZhName", "Mult", "Xyd", "Mat", "ColorR", "ColorG", "ColorB", "Price", "ProductKey", "Icon"],
		["string", "string", "number", "number", "string", "number", "number", "number", "number", "string", "string"],
		rows,
	)
	_save(wb, "AWMaterial.xlsx")


def gen_arm():
	wb = Workbook()
	rows = [
		[1, "新手之臂", 1, 1, 0, 0, "", 0, 234, 184, 146, "rbxassetid://74809006284031"],
		[2, "结实之臂", 1.5, 1, 75, 0, "", 0, 205, 150, 110, "rbxassetid://79060817271352"],
		[3, "青铜之臂", 2, 2, 250, 0, "", 0, 205, 127, 50, "rbxassetid://78998063862787"],
		[4, "钢铁之臂", 3, 2, 600, 0, "", 0, 160, 165, 175, "rbxassetid://125017192128258"],
		[5, "冰晶之臂", 4.5, 3, 1500, 0, "", 0, 110, 200, 255, "rbxassetid://135758805551826"],
		[6, "雷霆之臂", 6.5, 3, 4000, 0, "", 0, 255, 230, 60, "rbxassetid://98827391711336"],
		[7, "翡翠之臂", 9, 4, 10000, 0, "", 0, 60, 220, 110, "rbxassetid://104373057742300"],
		[8, "紫晶之臂", 13, 4, 40000, 0, "", 0, 170, 80, 255, "rbxassetid://116669000221342"],
		[9, "烈焰之臂", 18, 4, 150000, 0, "", 0, 255, 110, 30, "rbxassetid://125285035215604"],
		[10, "深渊之臂", 25, 4, 600000, 0, "", 0, 200, 30, 60, "rbxassetid://82645023916736"],
		[11, "神圣之臂", 40, 5, 0, 249, "HolyArm", 0, 255, 215, 90, "rbxassetid://75627413068113"],
		[12, "审判之臂", 60, 5, 0, 399, "JudgementArm", 0, 30, 30, 40, "rbxassetid://106638998049834"],
	]
	_sheet(
		wb,
		["手臂Id", "中文名", "力量倍率", "稀有度档", "现金价", "Robux", "商品键", "限量", "色R", "色G", "色B", "图标"],
		["Id", "ZhName", "Mult", "Xyd", "Price", "Robux", "ProductKey", "Limited", "ColorR", "ColorG", "ColorB", "Icon"],
		["number", "string", "number", "number", "number", "number", "string", "number", "number", "number", "number", "string"],
		rows,
	)
	_save(wb, "AWArm.xlsx")


def gen_dumbbell():
	wb = Workbook()
	rows = [
		[1, "训练哑铃", 1, 1, 0, 35, 35, 40, "Metal", 1, 190, 190, 200, "rbxassetid://79149397540649"],
		[2, "铁片哑铃", 1.3, 1, 150, 110, 110, 120, "DiamondPlate", 1.05, 190, 190, 200, "rbxassetid://108608298227670"],
		[3, "青铜哑铃", 1.7, 2, 600, 205, 127, 50, "Metal", 1.1, 150, 95, 45, "rbxassetid://91634863308487"],
		[4, "白银哑铃", 2.2, 2, 2000, 215, 220, 230, "Metal", 1.15, 160, 165, 175, "rbxassetid://83444554423906"],
		[5, "黄金哑铃", 3, 3, 7000, 255, 200, 40, "Metal", 1.2, 200, 150, 30, "rbxassetid://134229591004130"],
		[6, "冰晶哑铃", 4, 3, 25000, 120, 210, 255, "Ice", 1.25, 210, 240, 255, "rbxassetid://94366573381853"],
		[7, "翡翠哑铃", 5.5, 4, 90000, 40, 210, 110, "Glass", 1.3, 30, 120, 70, "rbxassetid://119581829333036"],
		[8, "紫晶哑铃", 7.5, 4, 350000, 170, 80, 255, "Glass", 1.35, 90, 40, 140, "rbxassetid://99369701283844"],
		[9, "熔岩哑铃", 10, 4, 1200000, 255, 100, 30, "CrackedLava", 1.4, 60, 40, 40, "rbxassetid://108797683466319"],
		[10, "钻石哑铃", 14, 4, 5000000, 120, 240, 255, "Neon", 1.5, 230, 250, 255, "rbxassetid://104387525439155"],
	]
	_sheet(
		wb,
		["哑铃Id", "中文名", "力量倍率", "稀有度档", "现金价", "片色R", "片色G", "片色B", "片材质", "片缩放", "柄色R", "柄色G", "柄色B", "图标"],
		["Id", "ZhName", "Mult", "Xyd", "Price", "PlateR", "PlateG", "PlateB", "PlateMat", "PlateScale", "HandleR", "HandleG", "HandleB", "Icon"],
		["number", "string", "number", "number", "number", "number", "number", "number", "string", "number", "number", "number", "number", "string"],
		rows,
	)
	_save(wb, "AWDumbbell.xlsx")


def gen_upgrade():
	wb = Workbook()
	rows = [
		[1, "TrainSpeed", 0, 1, 0],
		[2, "TrainSpeed", 1, 0.8, 50],
		[3, "TrainSpeed", 2, 0.6, 500],
		[4, "TrainSpeed", 3, 0.45, 6000],
		[5, "TrainSpeed", 4, 0.35, 80000],
		[6, "TrainSpeed", 5, 0.25, 900000],
		[7, "WalkSpeed", 0, 16, 0],
		[8, "WalkSpeed", 1, 20, 150],
		[9, "WalkSpeed", 2, 24, 1500],
		[10, "WalkSpeed", 3, 28, 20000],
		[11, "WalkSpeed", 4, 32, 250000],
		[12, "WalkSpeed", 5, 36, 2000000],
	]
	_sheet(
		wb,
		["行Id", "种类", "阶", "数值", "现金价"],
		["Id", "Kind", "Level", "Value", "Price"],
		["number", "string", "number", "number", "number"],
		rows,
	)
	_save(wb, "AWUpgrade.xlsx")


def gen_level():
	wb = Workbook()
	# PoolWeights 与 Pool 按位对齐；同材质可在不同箱有不同权重
	# 全关 Offer：抽取免费；购买价只在 AWMaterial.Price，本表不配价
	rows = [
		[1, "青草公园", "Grass", 0, "Grass", 100, 175, 70, "Offer", "材质木箱", "Sand|Wood|Grass", "60|48|30", "rbxassetid://138062702684407"],
		[2, "岩石峡谷", "Stone", 60, "Slate", 120, 125, 140, "Offer", "岩石石箱", "Wood|Grass|Snow|Rock", "48|30|22|12", "rbxassetid://108706762315626"],
		[3, "炙热沙漠", "Desert", 800, "Sand", 225, 200, 140, "Offer", "沙漠宝箱", "Snow|Rock|Metal|Water", "22|12|8|5", "rbxassetid://93020304140785"],
		[4, "冰雪山地", "Snow", 15000, "Snow", 235, 242, 250, "Offer", "冰晶宝箱", "Metal|Water|Brick|Magma", "8|5|3|1.5", "rbxassetid://136960778971634"],
		[5, "熔岩火山", "Lava", 250000, "Basalt", 55, 45, 45, "Offer", "熔岩宝箱", "Brick|Magma|Neon|Diamond", "3|1.5|0.8|0.4", "rbxassetid://127745719260945"],
	]
	_sheet(
		wb,
		["关卡", "中文名", "主题", "入场力量门槛", "地面材质", "地面R", "地面G", "地面B", "机台模式", "箱子名", "奖池材质Id", "奖池权重", "场景底图"],
		["Id", "ZhName", "Theme", "Gate", "GroundMat", "GroundR", "GroundG", "GroundB", "MachineMode", "MachineZhName", "Pool", "PoolWeights", "Icon"],
		["number", "string", "string", "number", "string", "number", "number", "number", "string", "string", "string[]", "number[]", "string"],
		rows,
	)
	_save(wb, "AWLevel.xlsx")


def gen_opponent():
	wb = Workbook()
	rows = [
		["L1_1", 1, "稻草人", "Scarecrow", 5, 25, 1],
		["L1_2", 1, "木头人", "Wooden", 15, 60, 1.1],
		["L1_3", 1, "树桩怪", "Stump", 40, 150, 1.25],
		["L2_1", 2, "小石怪", "Rock", 120, 400, 0.9],
		["L2_2", 2, "石头人", "Rock", 300, 900, 1.15],
		["L2_3", 2, "岩石巨人", "Rock", 700, 2000, 1.45],
		["L3_1", 3, "小仙人掌", "Cactus", 2000, 5000, 0.9],
		["L3_2", 3, "仙人掌战士", "Cactus", 5000, 12000, 1.15],
		["L3_3", 3, "仙人掌之王", "CactusKing", 12000, 28000, 1.4],
		["L4_1", 4, "小雪人", "Snowman", 35000, 60000, 0.9],
		["L4_2", 4, "大雪人", "Snowman", 90000, 140000, 1.2],
		["L4_3", 4, "雪怪", "Yeti", 220000, 320000, 1.4],
		["L5_1", 5, "熔岩小鬼", "Lava", 600000, 700000, 0.9],
		["L5_2", 5, "熔岩巨人", "Lava", 1500000, 1600000, 1.2],
		["L5_3", 5, "火山之王", "LavaKing", 4000000, 4000000, 1.5],
	]
	_sheet(
		wb,
		["对手Id", "关卡", "中文名", "模型", "力量", "现金奖励", "缩放"],
		["Id", "Level", "ZhName", "Model", "Power", "Reward", "Scale"],
		["string", "number", "string", "string", "number", "number", "number"],
		rows,
	)
	_save(wb, "AWEnemy.xlsx")


def main():
	OUT.mkdir(parents=True, exist_ok=True)
	gen_const()
	gen_material()
	gen_arm()
	gen_dumbbell()
	gen_upgrade()
	gen_level()
	gen_opponent()


if __name__ == "__main__":
	main()
