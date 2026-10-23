# -*- coding: utf-8 -*-
"""Rename Rarity -> Xyd (1..7) in AWMaterial / AWArm / AWDumbbell."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1] / "config"

MAP = {
	"Common": 1,
	"Uncommon": 2,
	"Rare": 3,
	"Epic": 4,
	"Legendary": 5,
	"Mythic": 6,
	"Exclusive": 7,
}


def patch(name: str):
	path = ROOT / name
	wb = load_workbook(path)
	ws = wb.active
	# row2 = field names
	col = None
	for c in range(1, ws.max_column + 1):
		if ws.cell(2, c).value == "Rarity":
			col = c
			break
	if col is None:
		raise SystemExit(f"{name}: Rarity column not found")
	ws.cell(1, col).value = "稀有度档(1普通..7专属)"
	ws.cell(2, col).value = "Xyd"
	ws.cell(3, col).value = "number"
	for r in range(4, ws.max_row + 1):
		raw = ws.cell(r, col).value
		if raw is None or raw == "":
			continue
		if isinstance(raw, (int, float)):
			ws.cell(r, col).value = int(raw)
			continue
		key = str(raw).strip()
		if key not in MAP:
			raise SystemExit(f"{name} row {r}: unknown Rarity {raw!r}")
		ws.cell(r, col).value = MAP[key]
	wb.save(path)
	print("patched", name, "col", col)


def main():
	for name in ("AWMaterial.xlsx", "AWArm.xlsx", "AWDumbbell.xlsx"):
		patch(name)


if __name__ == "__main__":
	main()
